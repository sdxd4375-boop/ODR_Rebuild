"""Unit tests for the extension modules (retrievers, docstore, sandbox,
charts, exporters, usage metering). No network and no real LLM required.

Run: uv run pytest tests/server -q
"""

import asyncio
import hashlib
import os
import subprocess

import pytest

from open_deep_research.docstore import LocalDocIndex, chunk_text
from open_deep_research.retrievers import (
    RETRIEVER_TOOLS,
    format_sources,
    load_extra_retrievers,
    parse_extra_retrievers,
)
from server.charts import inject_chart
from server.exporters import EXPORTERS, _parse_blocks
from server.routers.sessions import aggregate_usage
from server.sandbox import SandboxViolation, check_imports, run_python_code


def run(coro):
    return asyncio.run(coro)


# ---------------------------------------------------------------- retrievers
def test_retriever_registry_has_expected_engines():
    assert set(RETRIEVER_TOOLS) == {"duckduckgo", "arxiv", "local_docs"}


def test_parse_extra_retrievers_string_and_list():
    assert parse_extra_retrievers({"extra_retrievers": "duckduckgo, arxiv"}) == [
        "duckduckgo",
        "arxiv",
    ]
    assert parse_extra_retrievers({"extra_retrievers": ["local_docs"]}) == ["local_docs"]
    assert parse_extra_retrievers({}) == []
    assert parse_extra_retrievers({"extra_retrievers": None}) == []


def test_load_extra_retrievers_resolves_known_skips_unknown():
    config = {"configurable": {"extra_retrievers": "duckduckgo,bogus_engine"}}
    tools = run(load_extra_retrievers(config))
    assert [t.name for t in tools] == ["duckduckgo_web_search"]

    empty = run(load_extra_retrievers({"configurable": {}}))
    assert empty == []


def test_format_sources_matches_prompt_convention():
    out = format_sources(
        [{"title": "Paper", "url": "https://x", "content": "body text"}], "arXiv"
    )
    assert "--- SOURCE 1: Paper [arXiv] ---" in out
    assert "URL: https://x" in out
    assert "SUMMARY: body text" in out


def test_format_sources_truncates_long_content():
    out = format_sources([{"title": "t", "url": "u", "content": "x" * 99999}], "arXiv")
    assert "x" * 4001 not in out


# ---------------------------------------------------------------- docstore
def _fake_embed(texts):
    """Deterministic hash vectors — offline stand-in for real embeddings."""
    return [
        [float(b) / 255.0 for b in hashlib.md5(t.encode()).digest()[:4]]
        for t in texts
    ]


def test_chunk_text_sliding_window_and_pages():
    chunks = chunk_text("a" * 1200, size=500, overlap=50)
    assert len(chunks) == 3  # ceil((1200-500)/450) + 1
    assert all(len(c) <= 500 for c in chunks)

    assert len(chunk_text("page one\fpage two", size=500, overlap=50)) == 2


def test_docstore_add_search_remove_roundtrip(tmp_path):
    index = LocalDocIndex(str(tmp_path), embed_fn=_fake_embed)
    text = ("Revenue grew 20 percent in 2025. " * 40) + "\f" + ("Costs fell. " * 40)
    meta = index.add_document("report.pdf", text)

    doc_chunks = [c for c in index.chunks if c["doc_id"] == meta["doc_id"]]
    assert meta["chunks"] == len(doc_chunks)
    assert index.docs[meta["doc_id"]]["title"] == "report.pdf"

    hits = run(index.search("Revenue grew 20 percent", k=3))
    assert 0 < len(hits) <= 3
    assert all("title" in h and "page" in h for h in hits)

    # Persistence: a fresh instance over the same dir sees the same data.
    reloaded = LocalDocIndex(str(tmp_path), embed_fn=_fake_embed)
    assert reloaded.docs == index.docs

    assert reloaded.remove_document(meta["doc_id"]) is True
    assert not [c for c in reloaded.chunks if c["doc_id"] == meta["doc_id"]]
    assert reloaded.remove_document("nonexistent") is False


def test_docstore_search_empty_index(tmp_path):
    index = LocalDocIndex(str(tmp_path), embed_fn=_fake_embed)
    assert run(index.search("anything")) == []


# ---------------------------------------------------------------- sandbox
def test_check_imports_allows_whitelisted_blocks_rest():
    ok = "import numpy as np\nfrom statistics import mean\nimport matplotlib.pyplot as plt"
    assert check_imports(ok) == []

    flagged = check_imports("import os\nfrom subprocess import run\nimport socket")
    assert set(flagged) == {"os", "subprocess", "socket"}


def test_sandbox_runs_whitelisted_code():
    result = run_python_code("print(21 * 2)")
    assert result["stdout"].strip() == "42"


def test_sandbox_rejects_disallowed_import_without_executing():
    with pytest.raises(SandboxViolation) as exc_info:
        run_python_code("import socket\nprint('should not run')")
    assert "socket" in str(exc_info.value)


def test_sandbox_rejects_oversized_code():
    with pytest.raises(SandboxViolation):
        run_python_code("x = 1\n" * 40000)


def test_sandbox_timeout():
    with pytest.raises(subprocess.TimeoutExpired):
        run_python_code("while True:\n    pass", timeout_seconds=2)


def test_sandbox_can_render_matplotlib_chart(tmp_path):
    code = (
        "import matplotlib.pyplot as plt\n"
        "plt.bar(['a', 'b'], [1, 2])\n"
        "plt.savefig('chart_test.png', dpi=60)\n"
    )
    files = []
    try:
        from server.sandbox import run_python_code_collect_files

        files = run_python_code_collect_files(code, str(tmp_path))
    except Exception:  # noqa: BLE001 — headless env issues surface here
        pytest.fail("matplotlib chart rendering failed in sandbox")
    assert len(files) == 1 and files[0].endswith(".png")
    assert os.path.getsize(files[0]) > 0


# ---------------------------------------------------------------- charts
def test_inject_chart_after_first_section():
    md = "# Title\n\nintro\n\n## Section A\nbody\n"
    out = inject_chart(md, "Growth", "charts/c1.png")
    assert "![Growth](charts/c1.png)" in out
    assert out.index("![Growth]") < out.index("## Section A")


def test_inject_chart_appends_without_headings():
    out = inject_chart("just text", "T", "charts/c2.png")
    assert out.endswith("![T](charts/c2.png)\n")


# ---------------------------------------------------------------- exporters
SAMPLE_MD = """# Annual Report

Overview paragraph with findings.

## Numbers

- revenue grew 20%
- costs fell 5%

![Growth](charts/c1.png)
"""


def test_parse_blocks_structure():
    blocks = _parse_blocks(SAMPLE_MD)
    kinds = [b[0] for b in blocks]
    assert kinds[0] == "heading" and blocks[0][1] == "1"
    assert "image" in kinds
    assert "bullet" in kinds


def test_export_to_pptx_and_docx(tmp_path):
    sample_chart = tmp_path / "c1.png"
    sample_chart.write_bytes(b"\x89PNG fake")  # exporters only need the file to exist

    pptx_path = str(tmp_path / "out.pptx")
    EXPORTERS["pptx"](SAMPLE_MD, pptx_path, base_dir=str(tmp_path))
    assert os.path.getsize(pptx_path) > 0

    from pptx import Presentation

    prs = Presentation(pptx_path)
    titles = [s.shapes.title.text for s in prs.slides if s.shapes.title is not None]
    assert "Annual Report" in titles
    assert "Numbers" in titles

    docx_path = str(tmp_path / "out.docx")
    EXPORTERS["docx"](SAMPLE_MD, docx_path, base_dir=str(tmp_path))
    assert os.path.getsize(docx_path) > 0

    import docx as docx_lib

    text = "\n".join(p.text for p in docx_lib.Document(docx_path).paragraphs)
    assert "Annual Report" in text
    assert "revenue grew 20%" in text


# ---------------------------------------------------------------- usage metering
def test_aggregate_usage_last_chunk_per_call_wins():
    acc = {
        ("researcher", "m1"): {"input": 100, "output": 10, "total": 110, "model": "gpt"},
        # cumulative stream chunks for one call: only the final one is kept
        ("researcher", "m2"): {"input": 50, "output": 3, "total": 53, "model": "gpt"},
    }
    acc[("researcher", "m2")] = {"input": 50, "output": 9, "total": 59, "model": "gpt"}
    acc[("final_report_generation", "m3")] = {
        "input": 200, "output": 80, "total": 280, "model": "gpt",
    }

    agg = aggregate_usage(acc)
    assert agg["researcher"]["input"] == 150  # 100 + 50 (m2 counted once)
    assert agg["researcher"]["output"] == 19  # 10 + 9
    assert agg["researcher"]["calls"] == 2
    assert agg["final_report_generation"]["total"] == 280


def test_aggregate_usage_empty():
    assert aggregate_usage({}) == {}
